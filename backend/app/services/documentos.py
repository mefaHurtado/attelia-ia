"""Ingesta controlada de PDF, DOCX, TXT, MD y CSV en la colección personal."""
from io import BytesIO
from pathlib import Path
import hashlib
from langchain_core.documents import Document
from app.vectorstore.rag import dividir_documentos, obtener_vectorstore

EXTENSIONES = {".pdf", ".docx", ".txt", ".md", ".csv"}


def extraer_texto(nombre: str, contenido: bytes) -> str:
    ext = Path(nombre).suffix.lower()
    if ext not in EXTENSIONES:
        raise ValueError("Formato no compatible. Usa PDF, DOCX, TXT, MD o CSV.")
    if ext == ".pdf":
        from pypdf import PdfReader
        lector = PdfReader(BytesIO(contenido))
        if len(lector.pages) > 150:
            raise ValueError("PDF demasiado extenso (máximo 150 páginas).")
        return "\n\n".join((p.extract_text() or "") for p in lector.pages).strip()
    if ext == ".docx":
        from docx import Document as WordDoc
        doc = WordDoc(BytesIO(contenido))
        partes = [p.text for p in doc.paragraphs]
        for tabla in doc.tables:
            for fila in tabla.rows:
                partes.append(" | ".join(c.text for c in fila.cells))
        return "\n".join(partes).strip()
    try:
        return contenido.decode("utf-8-sig").strip()
    except UnicodeDecodeError:
        raise ValueError("El archivo debe estar codificado en UTF-8.")


def indexar_archivo(usuario_id: int, documento_id: int, nombre: str, contenido: bytes) -> int:
    texto = extraer_texto(nombre, contenido)
    if len(texto) < 40:
        raise ValueError("No se extrajo suficiente texto. Los PDF escaneados requieren OCR.")
    if len(texto) > 1_000_000:
        raise ValueError("El documento supera el límite de un millón de caracteres.")
    documento = Document(page_content=texto, metadata={"source": nombre, "document_id": int(documento_id), "user_id": int(usuario_id)})
    fragmentos = dividir_documentos([documento])
    base = obtener_vectorstore(usuario_id)
    for offset in range(0, len(fragmentos), 30):
        lote = fragmentos[offset:offset+30]
        ids = [hashlib.sha256(f"{usuario_id}|{documento_id}|{offset+i}|{d.page_content}".encode()).hexdigest() for i,d in enumerate(lote)]
        base.add_documents(lote, ids=ids)
    return len(fragmentos)


def eliminar_vectores(usuario_id: int, documento_id: int):
    base = obtener_vectorstore(usuario_id)
    ids = base.get(where={"document_id": int(documento_id)}).get("ids", [])
    if ids:
        base.delete(ids=ids)
